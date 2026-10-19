"""Create an independent PowerPoint test fixture, or validate a browser export."""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
if sys.argv[1]=='create':
    p=Presentation();p.slide_width=Inches(10);p.slide_height=Inches(7.5)
    s=p.slides.add_slide(p.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor(245,240,250)
    box=s.shapes.add_textbox(0,0,Inches(6),Inches(1));box.text='Independent PPTX';box.text_frame.paragraphs[0].runs[0].font.size=Pt(30)
    shape=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(1),Inches(2),Inches(2),Inches(1));shape.fill.solid();shape.fill.fore_color.rgb=RGBColor(130,95,180)
    s.notes_slide.notes_text_frame.text='Speaker notes from another application.'
    s2=p.slides.add_slide(p.slide_layouts[6]);s2.shapes.add_textbox(Inches(1),Inches(1),Inches(5),Inches(1)).text='Second imported slide'
    p.save(sys.argv[2]);print('Created independent PPTX fixture.')
else:
    p=Presentation(sys.argv[2]);assert len(p.slides)>=1;assert p.slide_width>0
    texts=[shape.text for slide in p.slides for shape in slide.shapes if shape.has_text_frame]
    assert any('Export proof' in t for t in texts),texts
    assert any('Keep these notes' in slide.notes_slide.notes_text_frame.text for slide in p.slides)
    print('PASS: python-pptx opens exported deck, editable text, notes and slide relationships.')
