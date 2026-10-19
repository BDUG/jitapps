"""Independent Office package checks; requires python-pptx and vsdx only for tests."""
import sys
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET
from pptx import Presentation
from pptx.util import Inches
from pptx.enum.shapes import MSO_CONNECTOR
from vsdx import VisioFile
import vsdx
if sys.argv[1]=='create':
    p=Presentation();s=p.slides.add_slide(p.slide_layouts[6])
    a=s.shapes.add_shape(1,Inches(1),Inches(1),Inches(2),Inches(1));a.text='Independent source'
    b=s.shapes.add_shape(1,Inches(5),Inches(1),Inches(2),Inches(1));b.text='Independent target'
    c=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(3),Inches(1.5),Inches(5),Inches(1.5));c.begin_connect(a,3);c.end_connect(b,1)
    p.save('/tmp/sketchix-independent.pptx')
    sample=Path(vsdx.__file__).parent/'media/media.vsdx'
    if not sample.exists():raise RuntimeError('Installed vsdx package has no independent fixture.')
    Path('/tmp/sketchix-independent.vsdx').write_bytes(sample.read_bytes())
    print('Created independent PPTX and Visio-library fixtures in /tmp.')
else:
    p=Presentation('/tmp/sketchix-export.pptx');assert len(p.slides)==1
    assert any(s.shape_type==9 for s in p.slides[0].shapes)
    assert any(s.has_text_frame and 'Start with an idea' in s.text for s in p.slides[0].shapes)
    with VisioFile('/tmp/sketchix-export.vsdx') as v:
        assert len(v.pages)==1 and len(v.pages[0].child_shapes)>=5
        assert any('Start with an idea' in s.text for s in v.pages[0].child_shapes)
    with ZipFile('/tmp/sketchix-export.vsdx') as z:
        root=ET.fromstring(z.read('visio/pages/page1.xml'));ns={'v':'http://schemas.microsoft.com/office/visio/2012/main'}
        assert len(root.findall('.//v:Connect',ns))==4
        ids={s.get('ID') for s in root.findall('.//v:Shape',ns)}
        for c in root.findall('.//v:Connect',ns):assert c.get('FromSheet') in ids and c.get('ToSheet') in ids
    print('PASS: independent libraries read editable PPTX/VSDX shapes; native connectors and Visio endpoint relationships are present.')
