"""Independent read check for browser-exported mind maps (development only)."""
import sys
from pptx import Presentation
p=Presentation(sys.argv[1]);assert len(p.slides)==1
shapes=p.slides[0].shapes
assert len([s for s in shapes if s.shape_type==9])>=1
assert any(s.has_text_frame and 'A world of ideas' in s.text for s in shapes)
assert any(s.has_text_frame and 'Delivery plan' in s.text for s in shapes)
print('PASS: python-pptx opens editable mind-map topics and native branch connectors.')
