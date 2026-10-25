"""Independent MSPDI/XLSX validation; development-only MPXJ/JPype/OpenPyXL."""
import sys
import mpxj
import jpype
jpype.startJVM()
try:
    from org.mpxj import ProjectFile, Duration, TimeUnit, RelationType, Relation
    from org.mpxj.mspdi import MSPDIWriter
    from org.mpxj.reader import UniversalProjectReader
    from java.time import LocalDateTime
    if sys.argv[1]=='create':
        p=ProjectFile();p.setDefaultCalendar(p.addDefaultBaseCalendar());props=p.getProjectProperties();props.setName('Independent project');props.setStartDate(LocalDateTime.of(2026,10,19,8,0));props.setFinishDate(LocalDateTime.of(2026,10,28,8,0))
        phase=p.addTask();phase.setName('Independent phase');phase.setSummary(True)
        a=phase.addTask();a.setName('Independent discovery');a.setDuration(Duration.getInstance(2,TimeUnit.DAYS));a.setStart(LocalDateTime.of(2026,10,19,8,0));a.setFinish(LocalDateTime.of(2026,10,20,17,0));a.setNotes('Independent source notes')
        b=phase.addTask();b.setName('Independent delivery');b.setDuration(Duration.getInstance(3,TimeUnit.DAYS));b.setStart(LocalDateTime.of(2026,10,22,8,0));b.setFinish(LocalDateTime.of(2026,10,26,17,0))
        builder=jpype.JClass('org.mpxj.Relation$Builder')();b.addPredecessor(builder.predecessorTask(a).successorTask(b).type(RelationType.FINISH_START).lag(Duration.getInstance(1,TimeUnit.DAYS)))
        resource=p.addResource();resource.setName('Independent owner');a.addResourceAssignment(resource)
        MSPDIWriter().write(p,'/tmp/projectix-independent.xml');print('Created independent MPXJ Project XML fixture.')
    else:
        p=UniversalProjectReader().read('/tmp/projectix-export.xml');tasks=[t for t in p.getTasks() if t.getName() is not None]
        assert any(str(t.getName())=='Integration task' for t in tasks)
        integration=next(t for t in tasks if str(t.getName())=='Integration task')
        assert len(integration.getPredecessors())==1
        assert any(str(t.getName())=='Launch milestone' and t.getMilestone() for t in tasks)
        assert any(str(r.getName())=='QA' for r in p.getResources())
        print('PASS: MPXJ reads exported Project XML tasks, native predecessor links, milestones and resources.')
        import openpyxl
        book=openpyxl.load_workbook('/tmp/projectix-export.xlsx');assert set(['Tasks','Dependencies','Project','Gantt']).issubset(book.sheetnames)
        rows=list(book['Tasks'].iter_rows(min_row=2,values_only=True));assert any(r[2]=='Integration task' for r in rows)
        assert book['Dependencies'].max_row>1;assert book['Tasks']['G2'].number_format=='yyyy-mm-dd'
        print('PASS: OpenPyXL reads task dates, dependencies, project settings and the Gantt worksheet.')
finally:
    jpype.shutdownJVM()
